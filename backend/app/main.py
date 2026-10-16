from pathlib import Path
from datetime import datetime
import io, csv, os
from typing import Optional
from fastapi import FastAPI, HTTPException, UploadFile, File, Form
from fastapi.responses import HTMLResponse, StreamingResponse
from fastapi.staticfiles import StaticFiles
from pydantic import BaseModel, Field
from sqlalchemy import create_engine, String, Integer, Float, DateTime, select, func
from sqlalchemy.orm import DeclarativeBase, Mapped, mapped_column, Session

BASE_DIR = Path(__file__).resolve().parent
DB_URL = os.getenv("DATABASE_URL", "sqlite:///./tt_erp.db")
connect_args = {"check_same_thread": False} if DB_URL.startswith("sqlite") else {}
engine = create_engine(DB_URL, connect_args=connect_args, future=True)

class Base(DeclarativeBase): pass
class Record(Base):
    __tablename__ = "records"
    id: Mapped[int] = mapped_column(Integer, primary_key=True, index=True)
    module: Mapped[str] = mapped_column(String(40), index=True)
    code: Mapped[str] = mapped_column(String(80), default="")
    name: Mapped[str] = mapped_column(String(180), index=True)
    details: Mapped[str] = mapped_column(String(500), default="")
    amount: Mapped[float] = mapped_column(Float, default=0)
    status: Mapped[str] = mapped_column(String(40), default="Active")
    created_at: Mapped[datetime] = mapped_column(DateTime, default=datetime.utcnow)

Base.metadata.create_all(engine)
app = FastAPI(title="TT ERP API", version="1.0.0", description="TT ERP local starter API")
app.mount("/static", StaticFiles(directory=str(BASE_DIR / "static")), name="static")

class RecordIn(BaseModel):
    module: str = Field(min_length=1, max_length=40)
    code: str = Field(default="", max_length=80)
    name: str = Field(min_length=1, max_length=180)
    details: str = Field(default="", max_length=500)
    amount: float = 0
    status: str = "Active"

@app.get("/", response_class=HTMLResponse)
def home():
    return (BASE_DIR / "static" / "index.html").read_text(encoding="utf-8")

@app.get("/api/health")
def health():
    with Session(engine) as db:
        db.execute(select(func.count()).select_from(Record)).scalar_one()
    return {"status":"ok", "app":"TT ERP", "database":"connected", "time":datetime.now().isoformat(timespec="seconds")}

@app.get("/api/dashboard")
def dashboard():
    with Session(engine) as db:
        total = db.scalar(select(func.count()).select_from(Record)) or 0
        customers = db.scalar(select(func.count()).select_from(Record).where(Record.module == "Customers")) or 0
        vendors = db.scalar(select(func.count()).select_from(Record).where(Record.module == "Vendors")) or 0
        sales = db.scalar(select(func.coalesce(func.sum(Record.amount), 0)).select_from(Record).where(Record.module == "Sales")) or 0
        recent = db.scalars(select(Record).order_by(Record.id.desc()).limit(8)).all()
        return {"total_records":total,"customers":customers,"vendors":vendors,"sales_total":round(float(sales),2),"recent":[serialize(x) for x in recent]}

def serialize(r):
    return {"id":r.id,"module":r.module,"code":r.code,"name":r.name,"details":r.details,"amount":r.amount,"status":r.status,"created_at":r.created_at.isoformat(timespec="seconds") if r.created_at else ""}

@app.get("/api/records")
def list_records(module: Optional[str] = None, q: Optional[str] = None):
    with Session(engine) as db:
        stmt = select(Record).order_by(Record.id.desc())
        if module: stmt = stmt.where(Record.module == module)
        if q: stmt = stmt.where(Record.name.ilike(f"%{q}%"))
        return [serialize(x) for x in db.scalars(stmt.limit(500)).all()]

@app.post("/api/records")
def create_record(payload: RecordIn):
    with Session(engine) as db:
        item = Record(**payload.model_dump())
        db.add(item); db.commit(); db.refresh(item)
        return serialize(item)

@app.put("/api/records/{record_id}")
def update_record(record_id: int, payload: RecordIn):
    with Session(engine) as db:
        item = db.get(Record, record_id)
        if not item: raise HTTPException(404, "Record not found")
        for k,v in payload.model_dump().items(): setattr(item,k,v)
        db.commit(); db.refresh(item); return serialize(item)

@app.delete("/api/records/{record_id}")
def delete_record(record_id: int):
    with Session(engine) as db:
        item = db.get(Record, record_id)
        if not item: raise HTTPException(404, "Record not found")
        db.delete(item); db.commit(); return {"ok":True,"deleted_id":record_id}

@app.get("/api/export.csv")
def export_csv(module: Optional[str] = None):
    with Session(engine) as db:
        stmt = select(Record).order_by(Record.id)
        if module: stmt = stmt.where(Record.module == module)
        records = [serialize(x) for x in db.scalars(stmt).all()]
    output = io.StringIO(); fields = ["id","module","code","name","details","amount","status","created_at"]
    writer = csv.DictWriter(output, fieldnames=fields); writer.writeheader(); writer.writerows(records)
    data = io.BytesIO(output.getvalue().encode("utf-8-sig"))
    return StreamingResponse(data, media_type="text/csv; charset=utf-8", headers={"Content-Disposition":"attachment; filename=tt_erp_export.csv"})

@app.get("/api/export.xlsx")
def export_xlsx(module: Optional[str] = None):
    from openpyxl import Workbook
    with Session(engine) as db:
        stmt = select(Record).order_by(Record.id)
        if module: stmt = stmt.where(Record.module == module)
        records = [serialize(x) for x in db.scalars(stmt).all()]
    wb = Workbook(); ws = wb.active; ws.title = "TT ERP Data"
    fields = ["id","module","code","name","details","amount","status","created_at"]
    ws.append(fields)
    for r in records: ws.append([r.get(f,"") for f in fields])
    ws.freeze_panes = "A2"; ws.auto_filter.ref = ws.dimensions
    for col in ws.columns:
        letter = col[0].column_letter
        ws.column_dimensions[letter].width = min(max(max(len(str(c.value or "")) for c in col)+2,12),38)
    buf = io.BytesIO(); wb.save(buf); buf.seek(0)
    return StreamingResponse(buf, media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet", headers={"Content-Disposition":"attachment; filename=tt_erp_export.xlsx"})

@app.post("/api/import/preview")
async def import_preview(file: UploadFile = File(...)):
    raw = await file.read()
    name = (file.filename or "").lower()
    rows = []
    try:
        if name.endswith(".xlsx"):
            from openpyxl import load_workbook
            wb = load_workbook(io.BytesIO(raw), read_only=True, data_only=True)
            ws = wb.active; it = ws.iter_rows(values_only=True)
            headers = [str(x or "").strip().lower() for x in next(it, [])]
            for row in it:
                rows.append(dict(zip(headers, row)))
        elif name.endswith(".csv"):
            text = raw.decode("utf-8-sig")
            rows = list(csv.DictReader(io.StringIO(text)))
        else: raise HTTPException(400, "Please upload a .csv or .xlsx file")
    except HTTPException: raise
    except Exception as exc: raise HTTPException(400, f"Could not read file: {exc}")
    return {"filename":file.filename,"rows_found":len(rows),"preview":rows[:10],"note":"Preview only. No data was imported or changed."}


@app.get("/api/reports/summary")
def report_summary():
    with Session(engine) as db:
        rows = db.execute(select(Record.module, func.count(Record.id), func.coalesce(func.sum(Record.amount), 0)).group_by(Record.module)).all()
        return {"modules":[{"module":m,"records":int(c),"total_amount":round(float(a or 0),2)} for m,c,a in rows]}

@app.post("/api/import/commit")
async def import_commit(file: UploadFile = File(...), module: str = Form("Customers")):
    allowed = {"Customers","Vendors","Inventory","Sales","Purchases","Production","Packing","DPR","Dispatch","Quality","Employees","Accounting","Maintenance","QR Labels"}
    if module not in allowed:
        raise HTTPException(400, "Unknown module")
    raw = await file.read()
    filename = (file.filename or "").lower()
    try:
        if filename.endswith(".xlsx"):
            from openpyxl import load_workbook
            wb = load_workbook(io.BytesIO(raw), read_only=True, data_only=True)
            ws = wb.active
            it = ws.iter_rows(values_only=True)
            headers = [str(x or "").strip().lower() for x in next(it, [])]
            rows = [dict(zip(headers, row)) for row in it]
        elif filename.endswith(".csv"):
            rows = list(csv.DictReader(io.StringIO(raw.decode("utf-8-sig"))))
        else:
            raise HTTPException(400, "Please upload a .csv or .xlsx file")
    except HTTPException:
        raise
    except Exception as exc:
        raise HTTPException(400, f"Could not read file: {exc}")
    imported = 0
    with Session(engine) as db:
        for row in rows:
            # Accept common header names; ignore rows without a name.
            norm = {str(k).strip().lower().replace(" ", "_"): v for k,v in row.items()}
            name = norm.get("name") or norm.get("description") or norm.get("item") or norm.get("customer_name") or norm.get("vendor_name")
            if name is None or not str(name).strip():
                continue
            try: amount = float(norm.get("amount") or norm.get("value") or 0)
            except (ValueError, TypeError): amount = 0.0
            item = Record(module=module, code=str(norm.get("code") or norm.get("id") or ""), name=str(name)[:180], details=str(norm.get("details") or norm.get("address") or norm.get("notes") or "")[:500], amount=amount, status=str(norm.get("status") or "Active")[:40])
            db.add(item); imported += 1
        db.commit()
    return {"ok":True,"module":module,"imported":imported,"rows_read":len(rows),"note":"Rows with a name/description were imported."}

@app.post("/api/diagnostics")
def diagnostics():
    checks = []
    checks.append({"name":"API process","status":"OK","message":"TT ERP API is responding."})
    try:
        with Session(engine) as db: db.execute(select(func.count()).select_from(Record)).scalar_one()
        checks.append({"name":"Database connection","status":"OK","message":"Local database is reachable."})
    except Exception as exc:
        checks.append({"name":"Database connection","status":"ERROR","message":str(exc)[:250]})
    checks.append({"name":"Excel export","status":"OK","message":"CSV and XLSX export endpoints are available."})
    return {"checks":checks,"checked_at":datetime.now().isoformat(timespec="seconds")}
